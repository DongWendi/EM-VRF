"""EM-VRF 10.3 precision correction. Executed in the original application globals."""
import types as _types

APP_VERSION = "2026.10.03-full-precision-r3"
QCH_DISPLAY_DECIMALS = 2


def canonical_delta_x_na(value):
    """Preserve the floating-point input; formatting belongs to the view."""
    return float(value)


def validate_delta_x_na(delta_x_na, x_na):
    raw_value = float(delta_x_na)
    initial_na = float(x_na)
    if not math.isfinite(raw_value):
        raise ValueError("ΔX_Na must be a finite number.")
    if raw_value < 0:
        raise ValueError(f"ΔX_Na={raw_value:.6g} is invalid. Required range: 0 ≤ ΔX_Na ≤ X_Na ({initial_na:.6g}). Prediction was blocked.")
    if raw_value > initial_na + NA_LIMIT_TOLERANCE:
        raise ValueError(f"ΔX_Na={raw_value:.6g} exceeds the initial Na stoichiometry X_Na={initial_na:.6g}. Required range: 0 ≤ ΔX_Na ≤ X_Na. Prediction was blocked.")
    return raw_value


def _sync_formula_from_dx(self, *args):
    try:
        value = self.formula_dx.value() * FARADAY / (self._formula_molar_mass() * 3.6)
        self._set_linked_value(self.formula_qch, value)
    except Exception:
        pass
    self._refresh_formula_na_guard()


def _sync_designer_from_dx(self, *args):
    try:
        value = self.des_dx.value() * FARADAY / (self._designer_molar_mass() * 3.6)
        self._set_linked_value(self.des_qch, value)
    except Exception:
        pass
    self._refresh_designer_na_guard()


def _export_dataframe_precision(frame, path):
    """XLSX retains numeric values with a two-decimal number format; CSV is raw."""
    if str(path).lower().endswith(".xlsx"):
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            frame.to_excel(writer, index=False)
            sheet = writer.sheets["Sheet1"]
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, (float, np.floating)):
                        cell.number_format = "0.00"
    else:
        frame.to_csv(path, index=False, encoding="utf-8-sig", float_format="%.17g")


def _export_method_factory(QFileDialog, QMessageBox):
    def batch_export_file(self):
        if self.batch_df is None or self.batch_df.empty:
            QMessageBox.information(self, "Export", "No batch results to export.")
            return
        path, selected_filter = QFileDialog.getSaveFileName(self, "Export results", "EM_VRF_predictions.xlsx", "Excel (*.xlsx);;CSV (*.csv)")
        if not path:
            return
        if Path(path).suffix.lower() not in (".xlsx", ".csv"):
            path += ".csv" if selected_filter.startswith("CSV") else ".xlsx"
        try:
            _export_dataframe_precision(self.batch_df, path)
        except Exception as exc:
            QMessageBox.critical(self, "Export failed", str(exc))
    return batch_export_file


_replacement_codes = {
    "_sync_formula_from_dx": _sync_formula_from_dx.__code__,
    "_sync_designer_from_dx": _sync_designer_from_dx.__code__,
    "batch_export_file": _export_method_factory(None, None).__code__,
}


def _rewrite_nested_precision(code):
    if code.co_name in _replacement_codes:
        replacement = _replacement_codes[code.co_name]
        if replacement.co_freevars != code.co_freevars:
            raise RuntimeError("Precision patch closure mismatch: " + code.co_name)
        return replacement
    constants = []
    for item in code.co_consts:
        if isinstance(item, _types.CodeType):
            item = _rewrite_nested_precision(item)
        elif code.co_name == "refresh" and isinstance(item, str):
            item = item.replace(" is used/displayed to ", " is displayed to ")
            item = item.replace("\n            decimals; ", "\n            decimals (full precision used internally); ")
            item = item.replace(" decimal;\n            ", " decimals;\n            ")
        constants.append(item)
    return code.replace(co_consts=tuple(constants))


_original_run_gui = _types.FunctionType(_rewrite_nested_precision(run_gui.__code__), globals(), "_original_run_gui", run_gui.__defaults__)


def _make_precision_spinbox(base):
    class FullPrecisionSpinBox(base):
        """Full floating-point value, compact idle display, full value when editing."""
        def __init__(self, *args, **kwargs):
            self._editing_precision = False
            self._edited_precision = False
            self._setting_precision = False
            super().__init__(*args, **kwargs)
            super().setDecimals(17)
            self._exact_precision = super().value()
            self.setKeyboardTracking(False)
            self.valueChanged.connect(self._capture_precision)
            self.lineEdit().textEdited.connect(self._mark_precision_edited)

        def setDecimals(self, decimals):
            super().setDecimals(17)

        def _mark_precision_edited(self, text):
            self._edited_precision = True

        def _capture_precision(self, value):
            self._qt_precision_signal = True
            if not self._setting_precision:
                self._exact_precision = float(value)
            self.setToolTip("Full precision: " + repr(self.value()) + "\nClick to edit the full value; the idle display uses two decimals.")

        def value(self):
            return getattr(self, "_exact_precision", super().value())

        def setValue(self, value):
            value = float(value)
            if not math.isfinite(value):
                raise ValueError("Input must be finite.")
            previous = self.value()
            self._exact_precision = min(max(value, self.minimum()), self.maximum())
            self._qt_precision_signal = False
            self._setting_precision = True
            try:
                super().setValue(self._exact_precision)
            finally:
                self._setting_precision = False
            if previous != self._exact_precision and not self._qt_precision_signal:
                # Qt may map two adjacent floating-point values to the same
                # 17-decimal storage value; linked controls still need a signal.
                self._setting_precision = True
                try:
                    self.valueChanged.emit(self._exact_precision)
                finally:
                    self._setting_precision = False
            self.setToolTip("Full precision: " + repr(self.value()) + "\nClick to edit the full value; the idle display uses two decimals.")

        def textFromValue(self, value):
            if getattr(self, "_editing_precision", False):
                return repr(float(value))
            return format(float(value), ".2f")

        def valueFromText(self, text):
            if not getattr(self, "_edited_precision", False) and hasattr(self, "_exact_precision"):
                return self._exact_precision
            return super().valueFromText(text)

        def focusInEvent(self, event):
            self._editing_precision = True
            self._edited_precision = False
            super().focusInEvent(event)
            self.lineEdit().setText(repr(self.value()))

        def focusOutEvent(self, event):
            super().focusOutEvent(event)
            self._editing_precision = False
            self._edited_precision = False
            self.lineEdit().setText(format(self.value(), ".2f"))

        def stepBy(self, steps):
            self.setValue(self.value() + int(steps) * self.singleStep())
    return FullPrecisionSpinBox


_PRECISION_CASES = [
    ("NMTS", "Na0.92Ni0.4Mn0.4Al0.05Ti0.1Li0.03Sb0.02O2", 149.24, -1.361105620589903),
    ("NMTS-A", "Na0.92Ni0.4Mn0.4Al0.08Ti0.08Li0.02Sb0.02O2", 139.2, -0.4934918803431927),
    ("NMTS-M", "Na0.905Ni0.4Mn0.4Al0.05Ti0.1Li0.015Mg0.02Mo0.015O2", 140.6, -0.6035408535903892),
    ("NMTS-L", "NaNi0.4Mn0.4Al0.06Ti0.04Li0.06Sb0.04O2", 150.2, -2.0087560879279645),
]


def self_test():
    assert ELEMENTS["Ce"].valence == 4 and ELEMENTS["Ce"].radius == 0.87
    ce_formula = "Na0.92Ni0.4Mn0.38Ce0.02Al0.05Ti0.1Li0.03Sb0.02O2"
    ce_descriptors = calculate_descriptors(parse_formula(ce_formula))
    assert abs(ce_descriptors["Ravg"] - 0.61685) < 1e-12
    expected_ce_phi = 5.325959421829658 + 0.02 * (4 / 0.87 - 4 / 0.53)
    assert abs(ce_descriptors["PhiTM"] - expected_ce_phi) < 1e-12
    results = []
    for name, formula, qch, expected in _PRECISION_CASES:
        result = predict_formula(formula, qch=qch)
        mass = calculate_descriptors(parse_formula(formula))["molar_mass"]
        raw_dx = qch * mass * 3.6 / FARADAY
        assert result["DeltaXNa"] == raw_dx
        assert abs(result["DeltaV_pred"] - expected) < 1e-12
        assert predict_formula(formula, delta_x_na=raw_dx)["DeltaV_pred"] == result["DeltaV_pred"]
        assert result["prediction_allowed"]
        results.append({"name": name, "Qch": qch, "DeltaXNa": raw_dx, "DeltaV_pred": result["DeltaV_pred"]})
    assert validate_delta_x_na(0.905, 0.905) == 0.905
    for value in (-0.1, float("nan"), float("inf"), 0.9051):
        try:
            validate_delta_x_na(value, 0.905)
        except ValueError:
            pass
        else:
            raise AssertionError("Invalid extraction was not blocked")
    mass = calculate_descriptors(parse_formula(_PRECISION_CASES[0][1]))["molar_mass"]
    try:
        predict_formula(_PRECISION_CASES[0][1], qch=0.93 * FARADAY / (mass * 3.6))
    except ValueError:
        pass
    else:
        raise AssertionError("Over-limit capacity was not blocked")
    phase_block = predict_formula("Na5/6Li5/18Mn13/18O2", delta_x_na=0.5)
    assert not phase_block["prediction_allowed"]
    return {"application_version": APP_VERSION, "passed": True, "ce_radius_A": ELEMENTS["Ce"].radius, "ce_descriptor_check": {"formula": ce_formula, "Ravg": ce_descriptors["Ravg"], "PhiTM": ce_descriptors["PhiTM"]}, "full_precision_cases": results, "physical_guards": "passed", "phase_gate": "passed"}


def _precision_gui_check(app, snapshot_path=None):
    from PySide6 import QtCore, QtGui, QtWidgets
    import tempfile
    import csv
    import openpyxl
    window = next(w for w in app.topLevelWidgets() if hasattr(w, "predict_tab"))
    tab = window.predict_tab
    assert tab.stack.currentIndex() == 1
    assert tab.formula_edit.text() == _PRECISION_CASES[0][1]
    assert tab.formula_mode_qch.isChecked() and not tab.formula_mode_dx.isChecked()
    assert tab.formula_qch.value() == 149.24
    assert tab.formula_dx.value() == predict_formula(_PRECISION_CASES[0][1], qch=149.24)["DeltaXNa"]
    startup_defaults = {"formula": tab.formula_edit.text(), "input_mode": "Qch", "Qch": tab.formula_qch.value(), "DeltaXNa": tab.formula_dx.value(), "formula_page_visible": True}
    assert "is used/displayed" not in window.model_info_tab.text.toPlainText()
    records = []
    for name, formula, qch, expected in _PRECISION_CASES:
        tab.formula_edit.setText(formula)
        tab.formula_mode_qch.setChecked(True)
        tab.formula_qch.setValue(qch)
        app.processEvents()
        assert tab.formula_qch.value() == qch
        raw_dx = predict_formula(formula, qch=qch)["DeltaXNa"]
        assert tab.formula_dx.value() == raw_dx
        tab.formula_dx.focusInEvent(QtGui.QFocusEvent(QtCore.QEvent.Type.FocusIn))
        tab.formula_dx.focusOutEvent(QtGui.QFocusEvent(QtCore.QEvent.Type.FocusOut))
        assert tab.formula_dx.value() == raw_dx
        assert tab.btn_formula.isEnabled()
        tab.run_formula()
        assert abs(tab.comparison_items[-1]["DeltaV_pred"] - expected) < 1e-12
        # A mode change uses the full linked value, even though idle text shows 2 decimals.
        tab.formula_mode_dx.setChecked(True)
        tab.run_formula()
        assert abs(tab.comparison_items[-1]["DeltaV_pred"] - expected) < 1e-12
        records.append({"name": name, "dx_internal": tab.formula_dx.value(), "dx_display": tab.formula_dx.textFromValue(tab.formula_dx.value()), "prediction": expected})
    # Commit a real text edit, and then blur/refocus without changing it.
    box = tab.formula_qch
    box.focusInEvent(QtGui.QFocusEvent(QtCore.QEvent.Type.FocusIn))
    box.lineEdit().setText("149.24")
    box.lineEdit().textEdited.emit("149.24")
    box.interpretText()
    box.focusOutEvent(QtGui.QFocusEvent(QtCore.QEvent.Type.FocusOut))
    assert box.value() == 149.24 and box.lineEdit().text() == "149.24"
    box.focusInEvent(QtGui.QFocusEvent(QtCore.QEvent.Type.FocusIn))
    box.focusOutEvent(QtGui.QFocusEvent(QtCore.QEvent.Type.FocusOut))
    assert box.value() == 149.24
    box.focusInEvent(QtGui.QFocusEvent(QtCore.QEvent.Type.FocusIn))
    box.lineEdit().setText("149.246789")
    box.lineEdit().textEdited.emit("149.246789")
    box.interpretText()
    box.focusOutEvent(QtGui.QFocusEvent(QtCore.QEvent.Type.FocusOut))
    assert box.value() == 149.246789 and box.lineEdit().text() == "149.25"
    small_values = []
    tiny_box = QtWidgets.QDoubleSpinBox()
    tiny_box.valueChanged.connect(small_values.append)
    tiny_box.setValue(1.0000000000000002e-5)
    tiny_box.setValue(1.0000000000000004e-5)
    assert len(small_values) == 2 and tiny_box.value() == 1.0000000000000004e-5
    # Element designer must retain three-decimal stoichiometries too.
    composition = parse_formula(_PRECISION_CASES[2][1])
    for element, checkbox in tab.selector.checkboxes.items():
        checkbox.setChecked(element in composition)
    for element, widget in tab.element_inputs.items():
        widget.setValue(composition.get(element, 0.0))
    tab.na_spin.setValue(composition["Na"])
    tab.o_spin.setValue(composition["O"])
    tab.des_mode_qch.setChecked(True)
    tab.des_qch.setValue(140.6)
    assert tab.na_spin.value() == 0.905
    assert tab.element_inputs["Li"].value() == 0.015
    assert tab.element_inputs["Mo"].value() == 0.015
    tab.run_designer()
    assert abs(tab.comparison_items[-1]["DeltaV_pred"] - _PRECISION_CASES[2][3]) < 1e-12
    tab.des_dx.setValue(0.906)
    tab.des_mode_dx.setChecked(True)
    assert not tab.btn_designer.isEnabled()
    tab.des_dx.setValue(0.905)
    assert tab.btn_designer.isEnabled()
    # Exercise the actual GUI batch loader and export methods with file dialogs
    # replaced only during this self-test; avoid interactive prompts.
    with tempfile.TemporaryDirectory(prefix="emvrf_precision_check_") as scratch:
        scratch = Path(scratch)
        fixture = scratch / "inputs.csv"
        pd.DataFrame([{"Name": n, "Formula": f, "Q_ch": q} for n, f, q, _ in _PRECISION_CASES]).to_csv(fixture, index=False)
        original_open = QtWidgets.QFileDialog.getOpenFileName
        original_save = QtWidgets.QFileDialog.getSaveFileName
        original_critical = QtWidgets.QMessageBox.critical
        def fail_dialog(*args):
            raise AssertionError("GUI error dialog: " + str(args[-1]))
        try:
            QtWidgets.QMessageBox.critical = staticmethod(fail_dialog)
            QtWidgets.QFileDialog.getOpenFileName = staticmethod(lambda *args: (str(fixture), "CSV (*.csv)"))
            tab.batch_load_file()
            assert len(tab.batch_df) == 4
            # Locate the numeric result by its values, preserving existing headers.
            delta_column = next(c for c in tab.batch_df.columns if all(isinstance(v, (int, float, np.floating)) and abs(float(v) - row[3]) < 1e-12 for v, row in zip(tab.batch_df[c], _PRECISION_CASES)))
            csv_path = scratch / "export.csv"
            QtWidgets.QFileDialog.getSaveFileName = staticmethod(lambda *args: (str(csv_path), "CSV (*.csv)"))
            tab.batch_export_file()
            with csv_path.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            assert all(float(row[delta_column]) == expected[3] for row, expected in zip(rows, _PRECISION_CASES))
            xlsx_base = scratch / "export_no_extension"
            QtWidgets.QFileDialog.getSaveFileName = staticmethod(lambda *args: (str(xlsx_base), "Excel (*.xlsx)"))
            tab.batch_export_file()
            workbook = openpyxl.load_workbook(str(xlsx_base) + ".xlsx", data_only=True)
            sheet = workbook.active
            column = next(cell.column for cell in sheet[1] if cell.value == delta_column)
            for index, expected in enumerate(_PRECISION_CASES, start=2):
                cell = sheet.cell(index, column)
                assert abs(cell.value - expected[3]) < 1e-14 and cell.number_format == "0.00"
            workbook.close()
        finally:
            QtWidgets.QFileDialog.getOpenFileName = original_open
            QtWidgets.QFileDialog.getSaveFileName = original_save
            QtWidgets.QMessageBox.critical = original_critical
    tab.formula_edit.setText(_PRECISION_CASES[0][1])
    tab.formula_mode_qch.setChecked(True)
    tab.formula_qch.setValue(149.24)
    tab.formula_edit.setFocus()
    app.processEvents()
    if snapshot_path:
        if not window.grab().save(str(snapshot_path)):
            raise RuntimeError("Unable to save GUI snapshot")
    result = {"application_version": APP_VERSION, "gui_passed": True, "startup_defaults": startup_defaults, "formula_and_mode_switch_cases": records, "text_edit_and_focus": "passed", "adjacent_float_signal": "passed", "designer_fractional_stoichiometry": "passed", "gui_physical_guard": "passed", "batch_import_and_csv_xlsx_export": "passed"}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


def run_gui(smoke_test=False, snapshot_path=None):
    from PySide6 import QtWidgets
    original_spin = QtWidgets.QDoubleSpinBox
    original_app = QtWidgets.QApplication
    QtWidgets.QDoubleSpinBox = _make_precision_spinbox(original_spin)
    class PrecisionApplication(original_app):
        def exec(self):
            window = next(w for w in self.topLevelWidgets() if hasattr(w, "predict_tab"))
            tab = window.predict_tab
            tab.formula_edit.setText(_PRECISION_CASES[0][1])
            tab.formula_mode_qch.setChecked(True)
            tab.formula_qch.setValue(149.24)
            tab._sync_formula_from_qch()
            tab.stack.setCurrentIndex(1)
            window.predict_tab._sync_designer_from_dx()
            if smoke_test or snapshot_path:
                result = _precision_gui_check(self, snapshot_path)
                if snapshot_path:
                    Path(str(snapshot_path) + ".json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
                for widget in self.topLevelWidgets():
                    widget.close()
                return 0
            return super().exec()
    QtWidgets.QApplication = PrecisionApplication
    try:
        return _original_run_gui(smoke_test=False, snapshot_path=None)
    finally:
        QtWidgets.QDoubleSpinBox = original_spin
        QtWidgets.QApplication = original_app
